from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Request, Depends, UploadFile, File, Form, Query, Header
from fastapi.responses import Response
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import logging
from pydantic import BaseModel, Field
from typing import List, Optional, Annotated
import uuid
import io
import json
import jwt
import bcrypt
import requests
from datetime import datetime, timezone, timedelta
import secrets
from fastapi import BackgroundTasks
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

# ----------------------------------------------------------------------------
# Setup
# ----------------------------------------------------------------------------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger("espak")

JWT_SECRET = os.environ["JWT_SECRET"]
JWT_ALGORITHM = "HS256"

# ----------------------------------------------------------------------------
# Object storage
# ----------------------------------------------------------------------------
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "espak"
storage_key = None

MIME_TYPES = {"pdf": "application/pdf", "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png"}
ALLOWED_EXT = {"pdf", "jpg", "jpeg", "png"}
MAX_FILE_SIZE = 2 * 1024 * 1024  # 2 MB


def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    resp.raise_for_status()
    return resp.json()


def get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")


# ----------------------------------------------------------------------------
# Auth helpers
# ----------------------------------------------------------------------------
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))
    except Exception:
        return False


def create_access_token(user_id: str, role: str) -> str:
    payload = {"sub": user_id, "role": role, "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "access"}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def clean(doc: dict) -> dict:
    if not doc:
        return doc
    doc = dict(doc)
    doc.pop("_id", None)
    doc.pop("password_hash", None)
    if doc.get("role") and not doc.get("roles"):
        doc["roles"] = [doc["role"]]
    if doc.get("roles") and not doc.get("role"):
        doc["role"] = doc["roles"][0]
    return doc


async def get_token_from_request(request: Request, auth: Optional[str] = None) -> Optional[str]:
    header = request.headers.get("Authorization", "")
    if header.startswith("Bearer "):
        return header[7:]
    if auth:
        return auth
    return None


async def get_current_user(request: Request) -> dict:
    header = request.headers.get("Authorization", "")
    token = header[7:] if header.startswith("Bearer ") else request.query_params.get("auth")
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        user = await db.users.find_one({"id": payload["sub"]})
        if not user:
            raise HTTPException(status_code=401, detail="Pengguna tidak ditemukan")
        return clean(user)
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sesi berakhir, silakan login kembali")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")


ALL_ROLES = ["admin", "kepala", "pegawai", "pj_program"]


def roles_of(user: dict) -> list:
    r = user.get("roles")
    if r:
        return list(r)
    return [user["role"]] if user.get("role") else []


def has_role(user: dict, *roles) -> bool:
    return bool(set(roles_of(user)) & set(roles))


def require_roles(*roles):
    async def checker(user: dict = Depends(get_current_user)) -> dict:
        if not (set(roles_of(user)) & set(roles)):
            raise HTTPException(status_code=403, detail="Akses ditolak untuk role ini")
        return user
    return checker


async def log_activity(user: dict, aktivitas: str, module: str, request: Request = None):
    ip = "-"
    if request:
        ip = request.client.host if request.client else "-"
    await db.audit_logs.insert_one({
        "id": str(uuid.uuid4()), "user_id": user.get("id"), "user_nama": user.get("nama"),
        "aktivitas": aktivitas, "module": module, "ip_address": ip, "created_at": now_iso(),
    })


async def notify(user_id: str, judul: str, pesan: str, tipe: str = "info"):
    await db.notifications.insert_one({
        "id": str(uuid.uuid4()), "user_id": user_id, "judul": judul, "pesan": pesan,
        "tipe": tipe, "is_read": False, "created_at": now_iso(),
    })


# ----------------------------------------------------------------------------
# Models
# ----------------------------------------------------------------------------
class LoginInput(BaseModel):
    identifier: str  # username or NIP
    password: str


class UserCreate(BaseModel):
    username: str
    nip: Optional[str] = ""
    password: str
    nama: str
    role: Optional[str] = None
    roles: Optional[List[str]] = None
    jabatan: Optional[str] = ""
    unit: Optional[str] = ""
    is_blud: Optional[bool] = False
    tipe_pegawai: Optional[str] = None


class UserUpdate(BaseModel):
    nama: Optional[str] = None
    nip: Optional[str] = None
    role: Optional[str] = None
    roles: Optional[List[str]] = None
    jabatan: Optional[str] = None
    unit: Optional[str] = None
    status: Optional[str] = None
    password: Optional[str] = None
    is_blud: Optional[bool] = None
    tipe_pegawai: Optional[str] = None


class VerifyInput(BaseModel):
    status: str  # disetujui | ditolak
    catatan: Optional[str] = ""


class ProgramInput(BaseModel):
    nama_program: str
    penanggung_jawab_id: Optional[str] = ""
    penanggung_jawab: Optional[str] = ""
    status: Optional[str] = "aktif"


class IndicatorInput(BaseModel):
    program_id: str
    nama_indikator: str
    target: float
    satuan: Optional[str] = "%"


class ReportInput(BaseModel):
    indicator_id: str
    bulan: int
    tahun: int
    numerator: float
    denominator: float
    target: float
    masalah: Optional[str] = ""
    analisis: Optional[str] = ""
    tindak_lanjut: Optional[str] = ""


class PolicyBriefInput(BaseModel):
    periode: str
    masalah: str
    analisis: str
    dampak: Optional[str] = ""
    rekomendasi: str
    penanggung_jawab: Optional[str] = ""
    target_penyelesaian: Optional[str] = ""
    status: Optional[str] = "Belum Ditindaklanjuti"


class SettingsInput(BaseModel):
    target_jpl: int
    target_sertifikat: int


# ----------------------------------------------------------------------------
# Computation helpers
# ----------------------------------------------------------------------------
async def get_settings() -> dict:
    s = await db.settings.find_one({"id": "global"})
    if not s:
        s = {"id": "global", "target_jpl": 40, "target_sertifikat": 8}
        await db.settings.insert_one(dict(s))
    return {"target_jpl": s["target_jpl"], "target_sertifikat": s["target_sertifikat"]}


def indicator_status(capaian: float, target: float) -> str:
    if target <= 0:
        return "hijau"
    ratio = (capaian / target) * 100
    if ratio >= 100:
        return "hijau"
    if ratio >= 80:
        return "kuning"
    return "merah"


def comp_status(jpl: float, cert: int, target_jpl: int, target_cert: int) -> str:
    if jpl == 0 and cert == 0:
        return "BELUM_MULAI"
    if jpl >= target_jpl and cert >= target_cert:
        return "MEMENUHI_JPL_DAN_SERTIFIKAT"
    if jpl >= target_jpl:
        return "MEMENUHI_JPL"
    return "DALAM_PROSES"


async def employee_stats(user: dict, settings: dict, tahun: Optional[int] = None) -> dict:
    q = {"employee_id": user["id"], "status": "disetujui"}
    if tahun:
        q["tahun"] = tahun
    certs = await db.certificates.find(q).to_list(1000)
    total_cert = len(certs)
    total_jpl = sum(c.get("jpl", 0) for c in certs)
    tj, ts = settings["target_jpl"], settings["target_sertifikat"]
    return {
        "id": user["id"], "nama": user["nama"], "nip": user.get("nip", ""),
        "jabatan": user.get("jabatan", ""), "unit": user.get("unit", ""),
        "role": (roles_of(user)[0] if roles_of(user) else ""), "roles": roles_of(user),
        "is_blud": user.get("is_blud", False),
        "total_jpl": total_jpl, "total_sertifikat": total_cert,
        "persen_jpl": round((total_jpl / tj) * 100) if tj else 0,
        "persen_sertifikat": round((total_cert / ts) * 100) if ts else 0,
        "kurang_jpl": max(0, tj - total_jpl), "kurang_sertifikat": max(0, ts - total_cert),
        "status": comp_status(total_jpl, total_cert, tj, ts),
    }


PERIODE_START = 2026


async def periode_options() -> list:
    """Available periods (years) for dashboards & uploads. Starts at 2026."""
    current_year = datetime.now(timezone.utc).year
    years = set(range(PERIODE_START, max(current_year, PERIODE_START) + 2))
    for y in await db.certificates.distinct("tahun"):
        try:
            y = int(y)
            if y >= PERIODE_START:
                years.add(y)
        except Exception:
            pass
    return sorted(years)


# ----------------------------------------------------------------------------
# Auth routes
# ----------------------------------------------------------------------------
@api_router.post("/auth/login")
async def login(data: LoginInput, request: Request):
    ident = data.identifier.strip()
    user = await db.users.find_one({"$or": [{"username": ident}, {"nip": ident}]})
    if not user or not verify_password(data.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Username/NIP atau password salah")
    if user.get("status") == "nonaktif":
        raise HTTPException(status_code=403, detail="Akun dinonaktifkan")
    token = create_access_token(user["id"], user["role"])
    await log_activity(clean(user), "Login ke sistem", "Auth", request)
    return {"token": token, "user": clean(user)}


@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user


@api_router.post("/auth/logout")
async def logout(user: dict = Depends(get_current_user)):
    return {"ok": True}


# ----------------------------------------------------------------------------
# Users management (admin)
# ----------------------------------------------------------------------------
@api_router.get("/users")
async def list_users(user: dict = Depends(require_roles("admin"))):
    users = await db.users.find().sort("created_at", -1).to_list(1000)
    return [clean(u) for u in users]


@api_router.post("/users")
async def create_user(data: UserCreate, request: Request, user: dict = Depends(require_roles("admin"))):
    exists = await db.users.find_one({"$or": [{"username": data.username}, {"nip": data.nip}] if data.nip else [{"username": data.username}]})
    if exists:
        raise HTTPException(status_code=400, detail="Username atau NIP sudah digunakan")
    roles = data.roles or ([data.role] if data.role else ["pegawai"])
    roles = [r for r in roles if r in ALL_ROLES] or ["pegawai"]
    tipe = data.tipe_pegawai or ("BLUD" if data.is_blud else "ASN")
    doc = {
        "id": str(uuid.uuid4()), "username": data.username, "nip": data.nip or "",
        "password_hash": hash_password(data.password), "nama": data.nama,
        "role": roles[0], "roles": roles, "is_blud": tipe == "BLUD", "tipe_pegawai": tipe,
        "jabatan": data.jabatan or "", "unit": data.unit or "", "status": "aktif",
        "created_at": now_iso(), "updated_at": now_iso(),
    }
    await db.users.insert_one(dict(doc))
    await log_activity(user, f"Menambah pengguna {data.nama}", "Manajemen Pengguna", request)
    return clean(doc)


@api_router.put("/users/{user_id}")
async def update_user(user_id: str, data: UserUpdate, request: Request, user: dict = Depends(require_roles("admin"))):
    update = {k: v for k, v in data.model_dump().items() if v is not None and k not in ("password", "roles", "role")}
    if data.roles is not None:
        roles = [r for r in data.roles if r in ALL_ROLES] or ["pegawai"]
        update["roles"] = roles
        update["role"] = roles[0]
    elif data.role is not None:
        update["role"] = data.role
        update["roles"] = [data.role]
    if data.password:
        update["password_hash"] = hash_password(data.password)
    if data.tipe_pegawai is not None:
        update["tipe_pegawai"] = data.tipe_pegawai
        update["is_blud"] = data.tipe_pegawai == "BLUD"
    elif "is_blud" in update:
        update["tipe_pegawai"] = "BLUD" if update["is_blud"] else "ASN"
    update["updated_at"] = now_iso()
    res = await db.users.update_one({"id": user_id}, {"$set": update})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    await log_activity(user, f"Memperbarui pengguna {user_id}", "Manajemen Pengguna", request)
    updated = await db.users.find_one({"id": user_id})
    return clean(updated)


@api_router.delete("/users/{user_id}")
async def delete_user(user_id: str, request: Request, user: dict = Depends(require_roles("admin"))):
    await db.users.delete_one({"id": user_id})
    await log_activity(user, f"Menghapus pengguna {user_id}", "Manajemen Pengguna", request)
    return {"ok": True}


# ----------------------------------------------------------------------------
# Employees monitoring
# ----------------------------------------------------------------------------
@api_router.get("/employees")
async def list_employees(user: dict = Depends(require_roles("admin", "kepala", "pj_program"))):
    settings = await get_settings()
    staff = await db.users.find({"roles": {"$in": ["pegawai", "pj_program"]}}).to_list(1000)
    result = [await employee_stats(s, settings) for s in staff]
    result.sort(key=lambda x: x["total_jpl"], reverse=True)
    return {"settings": settings, "employees": result}


@api_router.get("/me/stats")
async def my_stats(tahun: Optional[int] = None, user: dict = Depends(get_current_user)):
    settings = await get_settings()
    stats = await employee_stats(user, settings, tahun)
    return {"settings": settings, "stats": stats, "tahun": tahun, "periode_options": await periode_options()}


# ----------------------------------------------------------------------------
# Certificates
# ----------------------------------------------------------------------------
@api_router.get("/certificates")
async def list_certificates(
    employee_id: Optional[str] = None,
    status: Optional[str] = None,
    user: dict = Depends(get_current_user),
):
    q = {}
    if not has_role(user, "admin", "kepala"):
        q["employee_id"] = user["id"]
    elif employee_id:
        q["employee_id"] = employee_id
    if status:
        q["status"] = status
    certs = await db.certificates.find(q).sort("created_at", -1).to_list(1000)
    # attach employee name
    ids = list({c["employee_id"] for c in certs})
    users = await db.users.find({"id": {"$in": ids}}).to_list(1000)
    umap = {u["id"]: u for u in users}
    out = []
    for c in certs:
        c = clean(c)
        emp = umap.get(c["employee_id"], {})
        c["employee_nama"] = emp.get("nama", "-")
        c["employee_nip"] = emp.get("nip", "-")
        out.append(c)
    return out


@api_router.post("/certificates")
async def upload_certificate(
    request: Request,
    nama_pelatihan: str = Form(...),
    jenis_pelatihan: str = Form(""),
    penyelenggara: str = Form(""),
    nomor_sertifikat: str = Form(""),
    tanggal_pelatihan: str = Form(""),
    tahun: int = Form(...),
    jpl: float = Form(...),
    keterangan: str = Form(""),
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
):
    if tahun < PERIODE_START:
        raise HTTPException(status_code=400, detail=f"Tahun sertifikat minimal {PERIODE_START}.")
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "").lower()
    if ext not in ALLOWED_EXT:
        raise HTTPException(status_code=400, detail="Format file tidak diizinkan. Gunakan PDF, JPG, JPEG, atau PNG.")
    data = await file.read()
    if len(data) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="Ukuran file melebihi 2 MB. Silakan kompres file Anda.")
    content_type = MIME_TYPES.get(ext, "application/octet-stream")
    path = f"{APP_NAME}/certificates/{user['id']}/{uuid.uuid4()}.{ext}"
    result = put_object(path, data, content_type)
    doc = {
        "id": str(uuid.uuid4()), "employee_id": user["id"], "nama_pelatihan": nama_pelatihan,
        "jenis_pelatihan": jenis_pelatihan, "penyelenggara": penyelenggara, "nomor_sertifikat": nomor_sertifikat,
        "tanggal_pelatihan": tanggal_pelatihan, "tahun": tahun, "jpl": jpl, "keterangan": keterangan,
        "storage_path": result["path"], "original_filename": file.filename, "content_type": content_type,
        "file_size": len(data), "status": "menunggu", "catatan_verifikasi": "",
        "verified_by": "", "verified_at": "", "created_at": now_iso(),
    }
    await db.certificates.insert_one(dict(doc))
    await log_activity(user, f"Upload sertifikat: {nama_pelatihan}", "Sertifikat", request)
    await notify(user["id"], "Sertifikat Terunggah", f"Sertifikat '{nama_pelatihan}' berhasil diunggah dan menunggu verifikasi.", "info")
    return clean(doc)


@api_router.put("/certificates/{cert_id}/verify")
async def verify_certificate(cert_id: str, data: VerifyInput, request: Request, user: dict = Depends(require_roles("admin"))):
    if data.status not in ("disetujui", "ditolak"):
        raise HTTPException(status_code=400, detail="Status tidak valid")
    if data.status == "ditolak" and not (data.catatan or "").strip():
        raise HTTPException(status_code=400, detail="Alasan penolakan wajib diisi")
    cert = await db.certificates.find_one({"id": cert_id})
    if not cert:
        raise HTTPException(status_code=404, detail="Sertifikat tidak ditemukan")
    await db.certificates.update_one({"id": cert_id}, {"$set": {
        "status": data.status, "catatan_verifikasi": data.catatan or "",
        "verified_by": user["nama"], "verified_at": now_iso(),
    }})
    await log_activity(user, f"Verifikasi sertifikat {cert['nama_pelatihan']} -> {data.status}", "Verifikasi", request)
    if data.status == "disetujui":
        await notify(cert["employee_id"], "Sertifikat Disetujui", f"Sertifikat '{cert['nama_pelatihan']}' telah disetujui ({cert['jpl']} JPL).", "success")
    else:
        await notify(cert["employee_id"], "Sertifikat Ditolak", f"Sertifikat '{cert['nama_pelatihan']}' ditolak. Alasan: {data.catatan}", "error")
    return {"ok": True}


@api_router.delete("/certificates/{cert_id}")
async def delete_certificate(cert_id: str, user: dict = Depends(get_current_user)):
    cert = await db.certificates.find_one({"id": cert_id})
    if not cert:
        raise HTTPException(status_code=404, detail="Tidak ditemukan")
    if not has_role(user, "admin") and cert["employee_id"] != user["id"]:
        raise HTTPException(status_code=403, detail="Akses ditolak")
    await db.certificates.delete_one({"id": cert_id})
    return {"ok": True}


@api_router.get("/certificates/{cert_id}/file")
async def download_certificate(cert_id: str, request: Request, auth: Optional[str] = Query(None)):
    header = request.headers.get("Authorization", "")
    token = header[7:] if header.startswith("Bearer ") else auth
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")
    cert = await db.certificates.find_one({"id": cert_id})
    if not cert or not cert.get("storage_path"):
        raise HTTPException(status_code=404, detail="File tidak ditemukan")
    data, ct = get_object(cert["storage_path"])
    return Response(content=data, media_type=cert.get("content_type", ct))


# ----------------------------------------------------------------------------
# Programs & Indicators
# ----------------------------------------------------------------------------
@api_router.get("/programs")
async def list_programs(user: dict = Depends(get_current_user)):
    programs = await db.programs.find().to_list(1000)
    return [clean(p) for p in programs]


@api_router.post("/programs")
async def create_program(data: ProgramInput, request: Request, user: dict = Depends(require_roles("admin"))):
    doc = {"id": str(uuid.uuid4()), **data.model_dump(), "created_at": now_iso()}
    await db.programs.insert_one(dict(doc))
    await log_activity(user, f"Tambah program {data.nama_program}", "Program", request)
    return clean(doc)


@api_router.put("/programs/{pid}")
async def update_program(pid: str, data: ProgramInput, user: dict = Depends(require_roles("admin"))):
    await db.programs.update_one({"id": pid}, {"$set": data.model_dump()})
    return {"ok": True}


@api_router.delete("/programs/{pid}")
async def delete_program(pid: str, user: dict = Depends(require_roles("admin"))):
    await db.programs.delete_one({"id": pid})
    return {"ok": True}


@api_router.get("/indicators")
async def list_indicators(program_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    q = {"program_id": program_id} if program_id else {}
    inds = await db.indicators.find(q).to_list(1000)
    programs = await db.programs.find().to_list(1000)
    pmap = {p["id"]: p["nama_program"] for p in programs}
    out = []
    for i in inds:
        i = clean(i)
        i["nama_program"] = pmap.get(i["program_id"], "-")
        out.append(i)
    return out


@api_router.post("/indicators")
async def create_indicator(data: IndicatorInput, request: Request, user: dict = Depends(require_roles("admin", "pj_program"))):
    doc = {"id": str(uuid.uuid4()), **data.model_dump(), "status": "aktif", "created_at": now_iso()}
    await db.indicators.insert_one(dict(doc))
    await log_activity(user, f"Tambah indikator {data.nama_indikator}", "Indikator", request)
    return clean(doc)


@api_router.put("/indicators/{iid}")
async def update_indicator(iid: str, data: IndicatorInput, user: dict = Depends(require_roles("admin", "pj_program"))):
    await db.indicators.update_one({"id": iid}, {"$set": data.model_dump()})
    return {"ok": True}


@api_router.delete("/indicators/{iid}")
async def delete_indicator(iid: str, user: dict = Depends(require_roles("admin"))):
    await db.indicators.delete_one({"id": iid})
    return {"ok": True}


# ----------------------------------------------------------------------------
# Indicator reports
# ----------------------------------------------------------------------------
@api_router.get("/reports")
async def list_reports(bulan: Optional[int] = None, tahun: Optional[int] = None, program_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    q = {}
    if bulan:
        q["bulan"] = bulan
    if tahun:
        q["tahun"] = tahun
    reports = await db.indicator_reports.find(q).to_list(5000)
    indicators = await db.indicators.find().to_list(1000)
    programs = await db.programs.find().to_list(1000)
    imap = {i["id"]: i for i in indicators}
    pmap = {p["id"]: p["nama_program"] for p in programs}
    out = []
    for r in reports:
        r = clean(r)
        ind = imap.get(r["indicator_id"], {})
        if program_id and ind.get("program_id") != program_id:
            continue
        r["nama_indikator"] = ind.get("nama_indikator", "-")
        r["program_id"] = ind.get("program_id", "")
        r["nama_program"] = pmap.get(ind.get("program_id"), "-")
        r["satuan"] = ind.get("satuan", "%")
        out.append(r)
    return out


@api_router.post("/reports")
async def create_report(data: ReportInput, request: Request, user: dict = Depends(require_roles("admin", "pj_program"))):
    capaian = round((data.numerator / data.denominator) * 100, 2) if data.denominator else 0
    status = indicator_status(capaian, data.target)
    existing = await db.indicator_reports.find_one({"indicator_id": data.indicator_id, "bulan": data.bulan, "tahun": data.tahun})
    doc = {
        "indicator_id": data.indicator_id, "bulan": data.bulan, "tahun": data.tahun,
        "numerator": data.numerator, "denominator": data.denominator, "capaian": capaian,
        "target": data.target, "status": status, "masalah": data.masalah, "analisis": data.analisis,
        "tindak_lanjut": data.tindak_lanjut, "penanggung_jawab": user["nama"], "updated_at": now_iso(),
    }
    if existing:
        await db.indicator_reports.update_one({"id": existing["id"]}, {"$set": doc})
        rid = existing["id"]
    else:
        doc["id"] = str(uuid.uuid4())
        doc["created_at"] = now_iso()
        await db.indicator_reports.insert_one(dict(doc))
        rid = doc["id"]
    await log_activity(user, f"Input capaian indikator ({capaian}%)", "Input Data SPM", request)
    return {"ok": True, "id": rid, "capaian": capaian, "status": status}


# ----------------------------------------------------------------------------
# Dashboard
# ----------------------------------------------------------------------------
@api_router.get("/dashboard/stats")
async def dashboard_stats(tahun: Optional[int] = None, user: dict = Depends(get_current_user)):
    settings = await get_settings()
    staff = await db.users.find({"roles": {"$in": ["pegawai", "pj_program"]}}).to_list(1000)
    stats = [await employee_stats(s, settings, tahun) for s in staff]
    total_pegawai = len(stats)
    memenuhi_jpl = sum(1 for s in stats if s["total_jpl"] >= settings["target_jpl"])
    total_sert = sum(s["total_sertifikat"] for s in stats)
    total_jpl = sum(s["total_jpl"] for s in stats)
    memenuhi_target = sum(1 for s in stats if s["status"] == "MEMENUHI_JPL_DAN_SERTIFIKAT")
    programs = await db.programs.count_documents({})
    indicators = await db.indicators.count_documents({})

    # Aggregate per periode (tahun) across all approved certificates (>= 2026)
    all_certs = await db.certificates.find({"status": "disetujui"}).to_list(5000)
    yearly = {}
    monthly = {}
    for c in all_certs:
        y = c.get("tahun")
        if not y:
            try:
                y = datetime.fromisoformat(c.get("tanggal_pelatihan", "")).year
            except Exception:
                y = None
        if y and y >= PERIODE_START:
            e = yearly.setdefault(int(y), {"tahun": int(y), "jpl": 0, "sertifikat": 0})
            e["jpl"] += c.get("jpl", 0)
            e["sertifikat"] += 1
        # monthly (respect selected period if any)
        if tahun and (not y or int(y) != tahun):
            continue
        tgl = c.get("tanggal_pelatihan", "")
        try:
            key = datetime.fromisoformat(tgl).strftime("%Y-%m")
        except Exception:
            key = "lainnya"
        monthly[key] = monthly.get(key, 0) + c.get("jpl", 0)
    jpl_yearly = [yearly[k] for k in sorted(yearly)]
    monthly_list = [{"bulan": k, "jpl": v} for k, v in sorted(monthly.items()) if k != "lainnya"]

    ranking = sorted(stats, key=lambda x: x["total_jpl"], reverse=True)[:10]
    avg_jpl = round(total_jpl / total_pegawai, 1) if total_pegawai else 0
    jpl_values = sorted([s["total_jpl"] for s in stats])
    median_jpl = jpl_values[len(jpl_values) // 2] if jpl_values else 0

    return {
        "settings": settings,
        "tahun": tahun,
        "periode_options": await periode_options(),
        "cards": {
            "total_pegawai": total_pegawai,
            "memenuhi_jpl": memenuhi_jpl,
            "belum_memenuhi_jpl": total_pegawai - memenuhi_jpl,
            "total_sertifikat": total_sert,
            "total_jpl": total_jpl,
            "persen_memenuhi_target": round((memenuhi_target / total_pegawai) * 100) if total_pegawai else 0,
            "total_program": programs,
            "total_indikator": indicators,
        },
        "jpl_monthly": monthly_list,
        "jpl_yearly": jpl_yearly,
        "ranking": ranking,
        "avg_jpl": avg_jpl,
        "median_jpl": median_jpl,
        "distribusi_status": {
            "belum_mulai": sum(1 for s in stats if s["status"] == "BELUM_MULAI"),
            "dalam_proses": sum(1 for s in stats if s["status"] == "DALAM_PROSES"),
            "memenuhi_jpl": sum(1 for s in stats if s["status"] == "MEMENUHI_JPL"),
            "memenuhi_semua": memenuhi_target,
        },
    }


@api_router.get("/spm/dashboard")
async def spm_dashboard(bulan: Optional[int] = None, tahun: Optional[int] = None, program_id: Optional[str] = None,
                        start_bulan: Optional[int] = None, start_tahun: Optional[int] = None,
                        end_bulan: Optional[int] = None, end_tahun: Optional[int] = None,
                        user: dict = Depends(get_current_user)):
    now = datetime.now(timezone.utc)
    is_range = all([start_bulan, start_tahun, end_bulan, end_tahun])

    if is_range:
        s_idx = start_tahun * 12 + start_bulan
        e_idx = end_tahun * 12 + end_bulan
        if s_idx > e_idx:
            s_idx, e_idx = e_idx, s_idx
        all_reports = await list_reports(program_id=program_id, user=user)
        inrange = [r for r in all_reports if s_idx <= (r["tahun"] * 12 + r["bulan"]) <= e_idx]
        agg = {}
        for r in inrange:
            a = agg.setdefault(r["indicator_id"], {"num": 0, "den": 0, "target": r["target"],
                                                    "nama_indikator": r["nama_indikator"], "nama_program": r["nama_program"],
                                                    "program_id": r.get("program_id", ""), "satuan": r.get("satuan", "%")})
            a["num"] += r["numerator"]
            a["den"] += r["denominator"]
            a["target"] = r["target"]
        reports = []
        for iid, a in agg.items():
            cap = round(a["num"] / a["den"] * 100, 2) if a["den"] else 0
            reports.append({"id": iid, "indicator_id": iid, "nama_indikator": a["nama_indikator"],
                            "nama_program": a["nama_program"], "program_id": a["program_id"], "satuan": a["satuan"],
                            "bulan": 0, "tahun": 0, "numerator": a["num"], "denominator": a["den"],
                            "capaian": cap, "target": a["target"], "status": indicator_status(cap, a["target"])})
        trend_src = inrange
    else:
        tahun = tahun or now.year
        reports = await list_reports(bulan=bulan, tahun=tahun, program_id=program_id, user=user)
        trend_src = await list_reports(tahun=tahun, program_id=program_id, user=user)

    hijau = sum(1 for r in reports if r["status"] == "hijau")
    kuning = sum(1 for r in reports if r["status"] == "kuning")
    merah = sum(1 for r in reports if r["status"] == "merah")
    total = len(reports)
    trend = {}
    for r in trend_src:
        trend.setdefault(r["tahun"] * 12 + r["bulan"], []).append(r["capaian"])
    trend_list = [{"bulan": ((k - 1) % 12) + 1, "rata_capaian": round(sum(v) / len(v), 1)} for k, v in sorted(trend.items())]
    top_masalah = sorted([r for r in reports if r["status"] != "hijau"], key=lambda x: x["capaian"])[:5]
    return {
        "summary": {"total": total, "hijau": hijau, "kuning": kuning, "merah": merah,
                    "persen_tercapai": round((hijau / total) * 100) if total else 0},
        "reports": reports,
        "trend": trend_list,
        "top_masalah": top_masalah,
        "is_range": is_range,
    }


# ----------------------------------------------------------------------------
# Early Warning System
# ----------------------------------------------------------------------------
@api_router.get("/ews")
async def early_warning(user: dict = Depends(require_roles("admin", "kepala", "pj_program"))):
    settings = await get_settings()
    now = datetime.now(timezone.utc)
    warnings = []
    # Red/yellow indicators (current month)
    reports = await list_reports(bulan=now.month, tahun=now.year, user=user)
    for r in reports:
        if r["status"] == "merah":
            selisih = round(r["target"] - r["capaian"], 1)
            warnings.append({"tipe": "spm", "level": "merah",
                             "judul": f"Indikator {r['nama_indikator']} di bawah target",
                             "pesan": f"Capaian {r['nama_indikator']} ({r['nama_program']}) sebesar {r['capaian']}%, {selisih}% di bawah target {r['target']}%."})
        elif r["status"] == "kuning":
            warnings.append({"tipe": "spm", "level": "kuning",
                             "judul": f"Indikator {r['nama_indikator']} perlu perhatian",
                             "pesan": f"Capaian {r['nama_indikator']} sebesar {r['capaian']}% mendekati batas target {r['target']}%."})
    # Employees below JPL target
    staff = await db.users.find({"roles": {"$in": ["pegawai", "pj_program"]}}).to_list(1000)
    belum = []
    belum_upload = []
    for s in staff:
        st = await employee_stats(s, settings)
        if st["total_jpl"] < settings["target_jpl"]:
            belum.append(st)
        certs = await db.certificates.count_documents({"employee_id": s["id"]})
        if certs == 0:
            belum_upload.append(st["nama"])
    if belum:
        warnings.append({"tipe": "pegawai", "level": "kuning",
                         "judul": f"{len(belum)} pegawai belum mencapai target {settings['target_jpl']} JPL",
                         "pesan": f"Terdapat {len(belum)} pegawai yang masih memerlukan tambahan JPL untuk mencapai target."})
    if belum_upload:
        warnings.append({"tipe": "data", "level": "merah",
                         "judul": f"{len(belum_upload)} pegawai belum mengupload sertifikat",
                         "pesan": f"Pegawai berikut belum mengunggah sertifikat: {', '.join(belum_upload[:8])}{'...' if len(belum_upload) > 8 else ''}."})
    return {"warnings": warnings, "counts": {
        "merah": sum(1 for w in warnings if w["level"] == "merah"),
        "kuning": sum(1 for w in warnings if w["level"] == "kuning"),
        "pegawai_belum": len(belum), "belum_upload": len(belum_upload),
    }}


# ----------------------------------------------------------------------------
# Policy briefs
# ----------------------------------------------------------------------------
@api_router.get("/policy-briefs")
async def list_policy_briefs(user: dict = Depends(get_current_user)):
    briefs = await db.policy_briefs.find().sort("created_at", -1).to_list(1000)
    return [clean(b) for b in briefs]


@api_router.post("/policy-briefs")
async def create_policy_brief(data: PolicyBriefInput, request: Request, user: dict = Depends(require_roles("admin", "kepala", "pj_program"))):
    doc = {"id": str(uuid.uuid4()), **data.model_dump(), "dibuat_oleh": user["nama"], "created_at": now_iso()}
    await db.policy_briefs.insert_one(dict(doc))
    await log_activity(user, f"Membuat policy brief {data.periode}", "Policy Brief", request)
    return clean(doc)


@api_router.post("/policy-briefs/generate")
async def generate_policy_brief(user: dict = Depends(require_roles("admin", "kepala", "pj_program"))):
    """Rule-based auto-generation from red/yellow indicators of current period."""
    now = datetime.now(timezone.utc)
    reports = await list_reports(tahun=now.year, user=user)
    problems = sorted([r for r in reports if r["status"] != "hijau"], key=lambda x: x["capaian"])[:5]
    if not problems:
        masalah = "Seluruh indikator SPM berada pada status tercapai (hijau)."
        analisis = "Tidak ada indikator kritis pada periode ini."
        rekomendasi = "Pertahankan kinerja dan lanjutkan monitoring rutin bulanan."
        pj = "-"
    else:
        masalah = "; ".join([f"{p['nama_indikator']} ({p['nama_program']}) capaian {p['capaian']}% dari target {p['target']}%" for p in problems])
        analisis = "Rendahnya capaian diduga disebabkan keterbatasan sumber daya, cakupan sasaran belum optimal, dan pelaporan yang belum lengkap pada program terkait."
        rekomendasi = "Perkuat koordinasi lintas program, tingkatkan kunjungan/penjaringan sasaran, dan lakukan evaluasi mingguan pada indikator berstatus merah."
        pj = ", ".join(sorted({p["nama_program"] for p in problems}))
    return {
        "periode": now.strftime("%m-%Y"),
        "masalah": masalah,
        "analisis": analisis,
        "dampak": "Berpotensi menurunkan capaian SPM tahunan dan mutu pelayanan kesehatan masyarakat.",
        "rekomendasi": rekomendasi,
        "penanggung_jawab": pj,
        "target_penyelesaian": (now + timedelta(days=30)).strftime("%d-%m-%Y"),
        "status": "Belum Ditindaklanjuti",
    }


@api_router.put("/policy-briefs/{bid}")
async def update_policy_brief(bid: str, data: PolicyBriefInput, user: dict = Depends(require_roles("admin", "kepala", "pj_program"))):
    await db.policy_briefs.update_one({"id": bid}, {"$set": data.model_dump()})
    return {"ok": True}


@api_router.delete("/policy-briefs/{bid}")
async def delete_policy_brief(bid: str, user: dict = Depends(require_roles("admin", "kepala"))):
    await db.policy_briefs.delete_one({"id": bid})
    return {"ok": True}


# ----------------------------------------------------------------------------
# Notifications
# ----------------------------------------------------------------------------
@api_router.get("/notifications")
async def list_notifications(user: dict = Depends(get_current_user)):
    notifs = await db.notifications.find({"user_id": user["id"]}).sort("created_at", -1).to_list(100)
    unread = sum(1 for n in notifs if not n.get("is_read"))
    return {"notifications": [clean(n) for n in notifs], "unread": unread}


@api_router.put("/notifications/{nid}/read")
async def read_notification(nid: str, user: dict = Depends(get_current_user)):
    await db.notifications.update_one({"id": nid, "user_id": user["id"]}, {"$set": {"is_read": True}})
    return {"ok": True}


@api_router.put("/notifications/read-all")
async def read_all_notifications(user: dict = Depends(get_current_user)):
    await db.notifications.update_many({"user_id": user["id"]}, {"$set": {"is_read": True}})
    return {"ok": True}


# ----------------------------------------------------------------------------
# Audit logs
# ----------------------------------------------------------------------------
@api_router.get("/audit-logs")
async def list_audit_logs(user: dict = Depends(require_roles("admin"))):
    logs = await db.audit_logs.find().sort("created_at", -1).to_list(300)
    return [clean(l) for l in logs]


# ----------------------------------------------------------------------------
# Settings
# ----------------------------------------------------------------------------
@api_router.get("/settings")
async def get_settings_route(user: dict = Depends(get_current_user)):
    return await get_settings()


@api_router.put("/settings")
async def update_settings(data: SettingsInput, request: Request, user: dict = Depends(require_roles("admin"))):
    await db.settings.update_one({"id": "global"}, {"$set": {"target_jpl": data.target_jpl, "target_sertifikat": data.target_sertifikat}}, upsert=True)
    await log_activity(user, f"Ubah target: {data.target_jpl} JPL, {data.target_sertifikat} sertifikat", "Pengaturan", request)
    return await get_settings()


# ----------------------------------------------------------------------------
# Export (CSV / Excel / PDF berkop resmi)
# ----------------------------------------------------------------------------
KOP = ["PEMERINTAH KOTA SAMARINDA", "DINAS KESEHATAN", "UPTD PUSKESMAS PALARAN",
       "Jl. Puskesmas No. 1, Kec. Palaran, Kota Samarinda, Kalimantan Timur"]


def _csv_bytes(headers, rows):
    lines = [",".join(str(h) for h in headers)]
    for r in rows:
        lines.append(",".join('"' + str(c).replace('"', '""') + '"' for c in r))
    return "\n".join(lines).encode("utf-8")


def _xlsx_bytes(title, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan"
    ws.append([KOP[2]])
    ws.append([title])
    ws.append([f"Dicetak: {datetime.now(timezone.utc).strftime('%d-%m-%Y %H:%M UTC')}"])
    ws.append([])
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_bytes(title, headers, rows):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=16 * mm, bottomMargin=14 * mm, leftMargin=12 * mm, rightMargin=12 * mm)
    styles = getSampleStyleSheet()
    kop_style = ParagraphStyle("kop", parent=styles["Normal"], alignment=TA_CENTER, fontSize=13, leading=15, fontName="Helvetica-Bold")
    sub_style = ParagraphStyle("sub", parent=styles["Normal"], alignment=TA_CENTER, fontSize=8, textColor=colors.HexColor("#475569"))
    title_style = ParagraphStyle("title", parent=styles["Normal"], alignment=TA_CENTER, fontSize=11, spaceBefore=8, spaceAfter=10, fontName="Helvetica-Bold")
    elems = [Paragraph(KOP[0], kop_style), Paragraph(KOP[1], kop_style), Paragraph(KOP[2], kop_style),
             Paragraph(KOP[3], sub_style), Spacer(1, 4), Paragraph("_" * 200, sub_style),
             Paragraph(title, title_style)]
    data = [list(headers)] + [[str(c) for c in r] for r in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0284C7")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elems.append(table)
    elems.append(Spacer(1, 10))
    elems.append(Paragraph(f"Dicetak otomatis oleh Sistem E-SPAK · {datetime.now(timezone.utc).strftime('%d-%m-%Y %H:%M UTC')}", sub_style))
    doc.build(elems)
    return buf.getvalue()


def _export_response(fmt, name, title, headers, rows):
    if fmt == "xlsx":
        return Response(content=_xlsx_bytes(title, headers, rows),
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f"attachment; filename={name}.xlsx"})
    if fmt == "pdf":
        return Response(content=_pdf_bytes(title, headers, rows), media_type="application/pdf",
                        headers={"Content-Disposition": f"attachment; filename={name}.pdf"})
    return Response(content=_csv_bytes(headers, rows), media_type="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={name}.csv"})


@api_router.get("/export/employees")
async def export_employees(format: str = Query("csv"), user: dict = Depends(require_roles("admin", "kepala"))):
    settings = await get_settings()
    staff = await db.users.find({"roles": {"$in": ["pegawai", "pj_program"]}}).to_list(1000)
    headers = ["No", "Nama", "NIP", "Jabatan", "Unit", "Total JPL", "Total Sertifikat", "Persen JPL", "Status"]
    rows = []
    for i, s in enumerate(staff, 1):
        st = await employee_stats(s, settings)
        rows.append([i, st['nama'], st['nip'], st['jabatan'], st['unit'], st['total_jpl'], st['total_sertifikat'], f"{st['persen_jpl']}%", st['status']])
    return _export_response(format, "rekap_jpl_pegawai", "REKAP JPL & KINERJA PEGAWAI", headers, rows)


@api_router.get("/export/certificates")
async def export_certificates(format: str = Query("csv"), user: dict = Depends(require_roles("admin", "kepala"))):
    certs = await db.certificates.find().sort("created_at", -1).to_list(5000)
    ids = list({c["employee_id"] for c in certs})
    users = await db.users.find({"id": {"$in": ids}}).to_list(1000)
    umap = {u["id"]: u for u in users}
    headers = ["No", "Pegawai", "NIP", "Pelatihan", "Penyelenggara", "JPL", "Tahun", "Tanggal", "Status"]
    rows = []
    for i, c in enumerate(certs, 1):
        emp = umap.get(c["employee_id"], {})
        rows.append([i, emp.get("nama", "-"), emp.get("nip", "-"), c["nama_pelatihan"], c.get("penyelenggara", ""), c["jpl"], c.get("tahun", ""), c.get("tanggal_pelatihan", ""), c["status"]])
    return _export_response(format, "rekap_sertifikat", "REKAP SERTIFIKAT PELATIHAN", headers, rows)


@api_router.get("/export/spm")
async def export_spm(format: str = Query("csv"), tahun: Optional[int] = None, user: dict = Depends(require_roles("admin", "kepala"))):
    now = datetime.now(timezone.utc)
    reports = await list_reports(tahun=tahun or now.year, user=user)
    headers = ["Program", "Indikator", "Bulan", "Tahun", "Numerator", "Denominator", "Capaian", "Target", "Status"]
    rows = [[r['nama_program'], r['nama_indikator'], r['bulan'], r['tahun'], r['numerator'], r['denominator'], f"{r['capaian']}%", f"{r['target']}%", r['status']] for r in reports]
    return _export_response(format, "rekap_capaian_spm", "REKAP CAPAIAN SPM", headers, rows)


# ----------------------------------------------------------------------------
# Analytics pegawai
# ----------------------------------------------------------------------------
@api_router.get("/analytics/employees")
async def analytics_employees(user: dict = Depends(require_roles("admin", "kepala"))):
    settings = await get_settings()
    staff = await db.users.find({"roles": {"$in": ["pegawai", "pj_program"]}}).to_list(1000)
    stats = [await employee_stats(s, settings) for s in staff]
    top = sorted(stats, key=lambda x: x["total_jpl"], reverse=True)[:10]
    bottom = sorted(stats, key=lambda x: x["total_jpl"])[:10]
    units = {}
    for s in stats:
        units.setdefault(s["unit"] or "Lainnya", []).append(s)
    per_unit = []
    for u, arr in units.items():
        n = len(arr)
        per_unit.append({
            "unit": u, "jumlah": n,
            "rata_jpl": round(sum(a["total_jpl"] for a in arr) / n, 1) if n else 0,
            "rata_sertifikat": round(sum(a["total_sertifikat"] for a in arr) / n, 1) if n else 0,
            "memenuhi": sum(1 for a in arr if a["status"] == "MEMENUHI_JPL_DAN_SERTIFIKAT"),
        })
    per_unit.sort(key=lambda x: x["rata_jpl"], reverse=True)
    all_jpl = [s["total_jpl"] for s in stats]
    all_jpl_sorted = sorted(all_jpl)
    median = all_jpl_sorted[len(all_jpl_sorted) // 2] if all_jpl_sorted else 0
    return {
        "settings": settings, "top": top, "bottom": bottom, "per_unit": per_unit,
        "rata_jpl": round(sum(all_jpl) / len(all_jpl), 1) if all_jpl else 0,
        "median_jpl": median,
        "persen_memenuhi": round(sum(1 for s in stats if s["status"] == "MEMENUHI_JPL_DAN_SERTIFIKAT") / len(stats) * 100) if stats else 0,
    }


# ----------------------------------------------------------------------------
# Dataset endpoints (siap-sambung Google Sheets / Looker Studio)
# JSON default (community JSON connector) atau ?format=csv untuk IMPORTDATA
# Auth via header Bearer atau ?auth=<token>
# ----------------------------------------------------------------------------
def _to_csv(rows):
    if not rows:
        return ""
    header = list(rows[0].keys())
    lines = [",".join(header)] + [",".join(str(r[h]) for h in header) for r in rows]
    return "\n".join(lines)


@api_router.get("/dataset/jpl")
async def dataset_jpl(format: Optional[str] = Query(None), user: dict = Depends(get_current_user)):
    settings = await get_settings()
    staff = await db.users.find({"roles": {"$in": ["pegawai", "pj_program"]}}).to_list(1000)
    rows = []
    for s in staff:
        st = await employee_stats(s, settings)
        rows.append({
            "nama": st["nama"], "nip": st["nip"], "jabatan": st["jabatan"], "unit": st["unit"],
            "total_jpl": st["total_jpl"], "total_sertifikat": st["total_sertifikat"],
            "persen_jpl": st["persen_jpl"], "status": st["status"],
            "target_jpl": settings["target_jpl"], "target_sertifikat": settings["target_sertifikat"],
        })
    if format == "csv":
        return Response(content=_to_csv(rows), media_type="text/csv")
    return {"generated_at": now_iso(), "count": len(rows), "data": rows}


@api_router.get("/dataset/spm")
async def dataset_spm(format: Optional[str] = Query(None), tahun: Optional[int] = None, user: dict = Depends(get_current_user)):
    now = datetime.now(timezone.utc)
    reports = await list_reports(tahun=tahun or now.year, user=user)
    rows = [{
        "program": r["nama_program"], "indikator": r["nama_indikator"], "bulan": r["bulan"],
        "tahun": r["tahun"], "numerator": r["numerator"], "denominator": r["denominator"],
        "capaian": r["capaian"], "target": r["target"], "status": r["status"],
    } for r in reports]
    if format == "csv":
        return Response(content=_to_csv(rows), media_type="text/csv")
    return {"generated_at": now_iso(), "count": len(rows), "data": rows}


# ----------------------------------------------------------------------------
# Dataset sync (cron harian) — jaga snapshot Looker Studio selalu terbaru
# ----------------------------------------------------------------------------
WEBHOOK_CRON_SECRET = os.environ.get("WEBHOOK_CRON_SECRET", "")


async def _compute_jpl_dataset():
    settings = await get_settings()
    staff = await db.users.find({"roles": {"$in": ["pegawai", "pj_program"]}}).to_list(1000)
    out = []
    for s in staff:
        st = await employee_stats(s, settings)
        out.append({"nama": st["nama"], "nip": st["nip"], "unit": st["unit"], "total_jpl": st["total_jpl"],
                    "total_sertifikat": st["total_sertifikat"], "persen_jpl": st["persen_jpl"], "status": st["status"]})
    return out


async def _compute_spm_dataset():
    now = datetime.now(timezone.utc)
    reports = await list_reports(tahun=now.year, user={"role": "admin"})
    return [{"program": r["nama_program"], "indikator": r["nama_indikator"], "bulan": r["bulan"], "tahun": r["tahun"],
             "capaian": r["capaian"], "target": r["target"], "status": r["status"]} for r in reports]


async def run_dataset_sync():
    jpl = await _compute_jpl_dataset()
    spm = await _compute_spm_dataset()
    ts = now_iso()
    await db.dataset_snapshots.update_one({"id": "jpl"}, {"$set": {"data": jpl, "generated_at": ts, "count": len(jpl)}}, upsert=True)
    await db.dataset_snapshots.update_one({"id": "spm"}, {"$set": {"data": spm, "generated_at": ts, "count": len(spm)}}, upsert=True)
    await run_sip_alerts()
    logger.info(f"Dataset sync done: {len(jpl)} jpl rows, {len(spm)} spm rows")


SIP_ALERT_DAYS = 180


async def run_sip_alerts():
    """Kirim notifikasi ke pegawai bila SIP (terverifikasi) tinggal <= 180 hari."""
    sips = await db.sip.find({"status_verifikasi": "disetujui"}).to_list(5000)
    sent = 0
    for s in sips:
        _, days = sip_masa_status(s.get("tanggal_berakhir", ""))
        if days is None:
            continue
        if 0 <= days <= SIP_ALERT_DAYS and not s.get("alert_180_sent"):
            await notify(s["employee_id"], "SIP Akan Kedaluwarsa",
                         f"SIP {s.get('profesi','')} No. {s.get('nomor_sip','')} akan berakhir dalam {days} hari ({s.get('tanggal_berakhir')}). Segera lakukan perpanjangan.", "error")
            await db.sip.update_one({"id": s["id"]}, {"$set": {"alert_180_sent": True}})
            sent += 1
    if sent:
        logger.info(f"SIP alerts terkirim: {sent}")
    return sent


@api_router.post("/sip/run-alerts")
async def sip_run_alerts(user: dict = Depends(require_roles("admin"))):
    sent = await run_sip_alerts()
    return {"ok": True, "sent": sent}


@api_router.post("/cron/sync-dataset")
async def cron_sync_dataset(request: Request, background: BackgroundTasks):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else ""
    if not WEBHOOK_CRON_SECRET or not secrets.compare_digest(token, WEBHOOK_CRON_SECRET):
        raise HTTPException(status_code=401, detail="Unauthorized")
    background.add_task(run_dataset_sync)
    return {"ok": True, "accepted": True}


@api_router.get("/dataset/sync-status")
async def dataset_sync_status(user: dict = Depends(get_current_user)):
    snaps = await db.dataset_snapshots.find().to_list(10)
    return {s["id"]: {"generated_at": s.get("generated_at"), "count": s.get("count", 0)} for s in snaps}


# ----------------------------------------------------------------------------
# Backup & Restore (admin only)
# ----------------------------------------------------------------------------
BACKUP_COLLECTIONS = ["users", "certificates", "programs", "indicators",
                      "indicator_reports", "policy_briefs", "settings",
                      "notifications", "audit_logs", "dataset_snapshots"]


def _strip_mongo_id(doc: dict) -> dict:
    d = dict(doc)
    d.pop("_id", None)
    return d


@api_router.get("/backup/export")
async def backup_export(request: Request, user: dict = Depends(require_roles("admin"))):
    data = {}
    for col in BACKUP_COLLECTIONS:
        docs = await db[col].find().to_list(100000)
        data[col] = [_strip_mongo_id(d) for d in docs]
    payload = {
        "app": "E-SPAK",
        "version": 1,
        "exported_at": now_iso(),
        "exported_by": user.get("nama"),
        "collections": data,
    }
    await log_activity(user, "Ekspor backup penuh (JSON)", "Backup", request)
    content = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
    fname = f"espak_backup_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M')}.json"
    return Response(content=content, media_type="application/json",
                    headers={"Content-Disposition": f"attachment; filename={fname}"})


@api_router.post("/backup/import")
async def backup_import(request: Request, file: UploadFile = File(...),
                        mode: str = Form("replace"),
                        user: dict = Depends(require_roles("admin"))):
    raw = await file.read()
    try:
        payload = json.loads(raw.decode("utf-8"))
    except Exception:
        raise HTTPException(status_code=400, detail="File backup tidak valid (bukan JSON)")
    collections = payload.get("collections")
    if not isinstance(collections, dict):
        raise HTTPException(status_code=400, detail="Struktur backup tidak dikenali")
    mode = "merge" if mode == "merge" else "replace"
    summary = {}
    for col, docs in collections.items():
        if col not in BACKUP_COLLECTIONS or not isinstance(docs, list):
            continue
        if mode == "replace":
            await db[col].delete_many({})
        count = 0
        for d in docs:
            d = _strip_mongo_id(d)
            if d.get("id"):
                await db[col].update_one({"id": d["id"]}, {"$set": d}, upsert=True)
            else:
                await db[col].insert_one(d)
            count += 1
        summary[col] = count
    await log_activity(user, f"Pulihkan backup (mode: {mode})", "Backup", request)
    return {"ok": True, "mode": mode, "restored": summary}


@api_router.get("/backup/excel")
async def backup_excel(user: dict = Depends(require_roles("admin"))):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pengguna"
    ws.append(["No", "Nama", "Username", "NIP", "Email", "Role", "Jabatan", "Unit", "Status"])
    users = await db.users.find().to_list(10000)
    for i, u in enumerate(users, 1):
        ws.append([i, u.get("nama", ""), u.get("username", ""), u.get("nip", ""), u.get("email", ""),
                   u.get("role", ""), u.get("jabatan", ""), u.get("unit", ""), u.get("status", "")])

    programs = await db.programs.find().to_list(10000)
    pmap = {p["id"]: p.get("nama_program", "") for p in programs}
    ws2 = wb.create_sheet("Program")
    ws2.append(["No", "Nama Program", "Penanggung Jawab", "Status"])
    for i, p in enumerate(programs, 1):
        ws2.append([i, p.get("nama_program", ""), p.get("penanggung_jawab", ""), p.get("status", "")])

    inds = await db.indicators.find().to_list(10000)
    imap = {}
    ws3 = wb.create_sheet("Indikator")
    ws3.append(["No", "Program", "Indikator", "Target", "Satuan", "Status"])
    for i, ind in enumerate(inds, 1):
        imap[ind["id"]] = (ind.get("nama_indikator", ""), pmap.get(ind.get("program_id"), "-"))
        ws3.append([i, pmap.get(ind.get("program_id"), "-"), ind.get("nama_indikator", ""),
                    ind.get("target", ""), ind.get("satuan", "%"), ind.get("status", "")])

    ws4 = wb.create_sheet("Capaian SPM")
    ws4.append(["No", "Program", "Indikator", "Bulan", "Tahun", "Numerator", "Denominator", "Capaian(%)", "Target(%)", "Status"])
    reports = await db.indicator_reports.find().to_list(100000)
    for i, r in enumerate(reports, 1):
        nm, prog = imap.get(r.get("indicator_id"), ("-", "-"))
        ws4.append([i, prog, nm, r.get("bulan", ""), r.get("tahun", ""), r.get("numerator", ""),
                    r.get("denominator", ""), r.get("capaian", ""), r.get("target", ""), r.get("status", "")])

    buf = io.BytesIO()
    wb.save(buf)
    fname = f"espak_data_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={fname}"})


@api_router.get("/backup/csv")
async def backup_csv(dataset: str = Query("users"), user: dict = Depends(require_roles("admin"))):
    if dataset == "users":
        users = await db.users.find().to_list(10000)
        headers = ["No", "Nama", "Username", "NIP", "Email", "Role", "Jabatan", "Unit", "Status"]
        rows = [[i, u.get("nama", ""), u.get("username", ""), u.get("nip", ""), u.get("email", ""),
                 u.get("role", ""), u.get("jabatan", ""), u.get("unit", ""), u.get("status", "")]
                for i, u in enumerate(users, 1)]
        name = "data_pengguna"
    elif dataset == "spm":
        programs = await db.programs.find().to_list(10000)
        pmap = {p["id"]: p.get("nama_program", "") for p in programs}
        inds = await db.indicators.find().to_list(10000)
        imap = {ind["id"]: (ind.get("nama_indikator", ""), pmap.get(ind.get("program_id"), "-")) for ind in inds}
        reports = await db.indicator_reports.find().to_list(100000)
        headers = ["No", "Program", "Indikator", "Bulan", "Tahun", "Numerator", "Denominator", "Capaian(%)", "Target(%)", "Status"]
        rows = []
        for i, r in enumerate(reports, 1):
            nm, prog = imap.get(r.get("indicator_id"), ("-", "-"))
            rows.append([i, prog, nm, r.get("bulan", ""), r.get("tahun", ""), r.get("numerator", ""),
                         r.get("denominator", ""), r.get("capaian", ""), r.get("target", ""), r.get("status", "")])
        name = "capaian_spm"
    else:
        raise HTTPException(status_code=400, detail="dataset harus 'users' atau 'spm'")
    return Response(content=_csv_bytes(headers, rows), media_type="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={name}.csv"})


# ----------------------------------------------------------------------------
# Cuti (Leave) — pengajuan, verifikasi, saldo N/N-1/N-2 & cuti bersama
# ----------------------------------------------------------------------------
JENIS_CUTI = ["Tahunan", "Sakit"]
BALANCE_ORDER = ["saldo_bersama", "saldo_n2", "saldo_n1", "saldo_n"]
DEFAULT_LEAVE_CONFIG = {"default_saldo_n": 12, "default_saldo_bersama": 0, "jenis_cuti": ["Tahunan", "Sakit"]}


class LeaveConfigInput(BaseModel):
    default_saldo_n: float = 12
    default_saldo_bersama: float = 0
    jenis_cuti: List[str] = ["Tahunan", "Sakit"]
    tahun: Optional[int] = None


async def get_leave_config() -> dict:
    c = await db.settings.find_one({"id": "leave"})
    if not c:
        c = {"id": "leave", **DEFAULT_LEAVE_CONFIG, "tahun": datetime.now(timezone.utc).year}
        await db.settings.insert_one(dict(c))
    return {"default_saldo_n": c.get("default_saldo_n", 12),
            "default_saldo_bersama": c.get("default_saldo_bersama", 0),
            "jenis_cuti": c.get("jenis_cuti") or ["Tahunan", "Sakit"],
            "tahun": c.get("tahun", datetime.now(timezone.utc).year)}


class LeaveInput(BaseModel):
    employee_id: Optional[str] = None
    jenis: str = "Tahunan"
    tanggal_mulai: str
    tanggal_selesai: str
    alasan: Optional[str] = ""
    alamat: Optional[str] = ""


class LeaveVerify(BaseModel):
    status: str
    catatan: Optional[str] = ""


class BalanceInput(BaseModel):
    tahun: Optional[int] = None
    saldo_n: float = 0
    saldo_n1: float = 0
    saldo_n2: float = 0
    saldo_bersama: float = 0


def count_days(start: str, end: str) -> int:
    try:
        d1 = datetime.fromisoformat(start[:10]).date()
        d2 = datetime.fromisoformat(end[:10]).date()
    except Exception:
        return 0
    if d2 < d1:
        return 0
    return (d2 - d1).days + 1


async def get_balance(emp_id: str) -> dict:
    b = await db.leave_balances.find_one({"employee_id": emp_id})
    if not b:
        cfg = await get_leave_config()
        b = {"id": str(uuid.uuid4()), "employee_id": emp_id, "tahun": cfg["tahun"] or datetime.now(timezone.utc).year,
             "saldo_n": cfg["default_saldo_n"], "saldo_n1": 0, "saldo_n2": 0,
             "saldo_bersama": cfg["default_saldo_bersama"], "updated_at": now_iso()}
        await db.leave_balances.insert_one(dict(b))
    return b


def apply_deduction(bal: dict, days: int):
    remaining = days
    breakdown = {}
    for k in BALANCE_ORDER:
        if remaining <= 0:
            break
        take = min(bal.get(k, 0), remaining)
        if take > 0:
            bal[k] = bal.get(k, 0) - take
            breakdown[k] = take
            remaining -= take
    return breakdown, remaining


async def _save_balance(emp_id, bal):
    await db.leave_balances.update_one({"employee_id": emp_id}, {"$set": {**{k: bal[k] for k in BALANCE_ORDER}, "updated_at": now_iso()}})


@api_router.get("/leave/balances")
async def leave_balances(user: dict = Depends(get_current_user)):
    year = datetime.now(timezone.utc).year
    approved = await db.leaves.find({"status": "disetujui"}).to_list(5000)
    used = {}
    for l in approved:
        try:
            ly = int(str(l.get("tanggal_mulai", ""))[:4])
        except Exception:
            ly = None
        if ly == year:
            used[l["employee_id"]] = used.get(l["employee_id"], 0) + l.get("jumlah_hari", 0)
    if has_role(user, "admin", "kepala"):
        staff = await db.users.find({"roles": {"$in": ["pegawai", "pj_program"]}, "is_blud": True}).to_list(1000)
        out = []
        for s in staff:
            b = await get_balance(s["id"])
            out.append({**clean(b), "nama": s["nama"], "nip": s.get("nip", ""),
                        "unit": s.get("unit", ""), "jabatan": s.get("jabatan", ""),
                        "is_blud": s.get("is_blud", False),
                        "tipe_pegawai": s.get("tipe_pegawai", "BLUD" if s.get("is_blud") else "ASN"),
                        "dipakai": used.get(s["id"], 0)})
        return out
    b = await get_balance(user["id"])
    return [{**clean(b), "nama": user["nama"], "nip": user.get("nip", ""),
             "unit": user.get("unit", ""), "jabatan": user.get("jabatan", ""),
             "is_blud": user.get("is_blud", False),
             "tipe_pegawai": user.get("tipe_pegawai", "BLUD" if user.get("is_blud") else "ASN"),
             "dipakai": used.get(user["id"], 0)}]


@api_router.get("/leave/my-balance")
async def my_leave_balance(user: dict = Depends(get_current_user)):
    year = datetime.now(timezone.utc).year
    approved = await db.leaves.find({"employee_id": user["id"], "status": "disetujui"}).to_list(2000)
    dipakai = sum(l.get("jumlah_hari", 0) for l in approved if str(l.get("tanggal_mulai", ""))[:4] == str(year))
    return {**clean(await get_balance(user["id"])), "dipakai": dipakai}


@api_router.put("/leave/balances/{employee_id}")
async def set_leave_balance(employee_id: str, data: BalanceInput, request: Request, user: dict = Depends(require_roles("admin"))):
    await get_balance(employee_id)
    upd = {"saldo_n": data.saldo_n, "saldo_n1": data.saldo_n1, "saldo_n2": data.saldo_n2,
           "saldo_bersama": data.saldo_bersama, "updated_at": now_iso()}
    if data.tahun:
        upd["tahun"] = data.tahun
    await db.leave_balances.update_one({"employee_id": employee_id}, {"$set": upd})
    await log_activity(user, f"Atur saldo cuti pegawai {employee_id}", "Cuti", request)
    return {"ok": True}


@api_router.get("/leave/config")
async def leave_config_get(user: dict = Depends(get_current_user)):
    return await get_leave_config()


@api_router.put("/leave/config")
async def leave_config_put(data: LeaveConfigInput, request: Request, user: dict = Depends(require_roles("admin"))):
    jenis = [j.strip() for j in data.jenis_cuti if j and j.strip()]
    if not jenis:
        jenis = ["Tahunan", "Sakit"]
    upd = {"default_saldo_n": data.default_saldo_n, "default_saldo_bersama": data.default_saldo_bersama,
           "jenis_cuti": jenis, "tahun": data.tahun or datetime.now(timezone.utc).year}
    await db.settings.update_one({"id": "leave"}, {"$set": upd}, upsert=True)
    await log_activity(user, "Ubah konfigurasi cuti pegawai", "Pengaturan", request)
    return await get_leave_config()


@api_router.post("/leave/config/apply")
async def leave_config_apply(request: Request, user: dict = Depends(require_roles("admin"))):
    cfg = await get_leave_config()
    staff = await db.users.find({"roles": {"$in": ["pegawai", "pj_program"]}, "is_blud": True}).to_list(1000)
    count = 0
    for s in staff:
        await get_balance(s["id"])
        await db.leave_balances.update_one({"employee_id": s["id"]}, {"$set": {
            "saldo_n": cfg["default_saldo_n"], "saldo_bersama": cfg["default_saldo_bersama"],
            "tahun": cfg["tahun"], "updated_at": now_iso()}})
        count += 1
    await log_activity(user, f"Terapkan konfigurasi cuti ke {count} pegawai BLUD", "Pengaturan", request)
    return {"ok": True, "updated": count}


async def _leave_with_names(leaves):
    ids = list({l["employee_id"] for l in leaves})
    users = await db.users.find({"id": {"$in": ids}}).to_list(1000)
    umap = {u["id"]: u for u in users}
    out = []
    for l in leaves:
        l = clean(l)
        emp = umap.get(l["employee_id"], {})
        l["employee_nama"] = emp.get("nama", "-")
        l["employee_nip"] = emp.get("nip", "-")
        l["employee_jabatan"] = emp.get("jabatan", "-")
        l["employee_unit"] = emp.get("unit", "-")
        l["employee_tipe"] = emp.get("tipe_pegawai", "BLUD" if emp.get("is_blud") else "ASN")
        l["has_lampiran"] = bool(l.get("lampiran_path"))
        out.append(l)
    return out


@api_router.get("/leaves")
async def list_leaves(status: Optional[str] = None, employee_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    q = {}
    if not has_role(user, "admin", "kepala"):
        q["employee_id"] = user["id"]
    elif employee_id:
        q["employee_id"] = employee_id
    if status:
        q["status"] = status
    leaves = await db.leaves.find(q).sort("created_at", -1).to_list(2000)
    return await _leave_with_names(leaves)


@api_router.post("/leaves")
async def create_leave(data: LeaveInput, request: Request, user: dict = Depends(get_current_user)):
    cfg = await get_leave_config()
    if data.jenis not in cfg["jenis_cuti"]:
        raise HTTPException(status_code=400, detail="Jenis cuti tidak valid")
    is_admin = has_role(user, "admin")
    on_behalf = is_admin and data.employee_id and data.employee_id != user["id"]
    emp_id = data.employee_id if (is_admin and data.employee_id) else user["id"]
    emp = await db.users.find_one({"id": emp_id})
    if not emp:
        raise HTTPException(status_code=404, detail="Pegawai tidak ditemukan")
    days = count_days(data.tanggal_mulai, data.tanggal_selesai)
    if days <= 0:
        raise HTTPException(status_code=400, detail="Rentang tanggal cuti tidak valid")
    doc = {
        "id": str(uuid.uuid4()), "employee_id": emp_id, "jenis": data.jenis,
        "tanggal_mulai": data.tanggal_mulai, "tanggal_selesai": data.tanggal_selesai,
        "jumlah_hari": days, "alasan": data.alasan or "", "alamat": data.alamat or "",
        "status": "diajukan", "catatan": "", "deduksi": {},
        "diajukan_oleh": user["nama"], "verified_by": "", "verified_at": "", "created_at": now_iso(),
    }
    if on_behalf:
        if data.jenis == "Tahunan":
            bal = await get_balance(emp_id)
            breakdown, remaining = apply_deduction(bal, days)
            if remaining > 0:
                raise HTTPException(status_code=400, detail=f"Saldo cuti tidak cukup. Kurang {remaining} hari.")
            await _save_balance(emp_id, bal)
            doc["deduksi"] = breakdown
        doc["status"] = "disetujui"
        doc["verified_by"] = user["nama"]
        doc["verified_at"] = now_iso()
    await db.leaves.insert_one(dict(doc))
    await log_activity(user, f"Pengajuan cuti {data.jenis} ({days} hari)", "Cuti", request)
    if doc["status"] == "diajukan":
        for a in await db.users.find({"roles": "admin"}).to_list(100):
            await notify(a["id"], "Pengajuan Cuti Baru", f"{emp['nama']} mengajukan cuti {data.jenis} {days} hari.", "info")
    else:
        await notify(emp_id, "Cuti Ditambahkan", f"Cuti {data.jenis} {days} hari telah dicatat & disetujui admin.", "success")
    return clean(doc)


@api_router.put("/leaves/{leave_id}/verify")
async def verify_leave(leave_id: str, data: LeaveVerify, request: Request, user: dict = Depends(require_roles("admin"))):
    if data.status not in ("disetujui", "ditolak"):
        raise HTTPException(status_code=400, detail="Status tidak valid")
    if data.status == "ditolak" and not (data.catatan or "").strip():
        raise HTTPException(status_code=400, detail="Alasan penolakan wajib diisi")
    lv = await db.leaves.find_one({"id": leave_id})
    if not lv:
        raise HTTPException(status_code=404, detail="Pengajuan cuti tidak ditemukan")
    if lv["status"] != "diajukan":
        raise HTTPException(status_code=400, detail="Pengajuan sudah diproses")
    upd = {"status": data.status, "catatan": data.catatan or "", "verified_by": user["nama"], "verified_at": now_iso()}
    if data.status == "disetujui" and lv["jenis"] == "Tahunan":
        bal = await get_balance(lv["employee_id"])
        breakdown, remaining = apply_deduction(bal, lv["jumlah_hari"])
        if remaining > 0:
            raise HTTPException(status_code=400, detail=f"Saldo cuti tidak cukup. Kurang {remaining} hari.")
        await _save_balance(lv["employee_id"], bal)
        upd["deduksi"] = breakdown
    await db.leaves.update_one({"id": leave_id}, {"$set": upd})
    await log_activity(user, f"Verifikasi cuti -> {data.status}", "Cuti", request)
    tipe = "success" if data.status == "disetujui" else "error"
    extra = f" Alasan: {data.catatan}" if data.status == "ditolak" else ""
    await notify(lv["employee_id"], f"Cuti {data.status.capitalize()}", f"Pengajuan cuti {lv['jenis']} {lv['jumlah_hari']} hari {data.status}.{extra}", tipe)
    return {"ok": True}


@api_router.put("/leaves/{leave_id}/cancel")
async def cancel_leave(leave_id: str, request: Request, user: dict = Depends(require_roles("admin"))):
    lv = await db.leaves.find_one({"id": leave_id})
    if not lv:
        raise HTTPException(status_code=404, detail="Tidak ditemukan")
    if lv["status"] == "dibatalkan":
        raise HTTPException(status_code=400, detail="Cuti sudah dibatalkan")
    if lv["status"] == "disetujui" and lv.get("deduksi"):
        bal = await get_balance(lv["employee_id"])
        for k, v in lv["deduksi"].items():
            bal[k] = bal.get(k, 0) + v
        await _save_balance(lv["employee_id"], bal)
    await db.leaves.update_one({"id": leave_id}, {"$set": {"status": "dibatalkan", "deduksi": {}, "catatan": "Dibatalkan oleh admin", "verified_by": user["nama"], "verified_at": now_iso()}})
    await log_activity(user, "Pembatalan cuti", "Cuti", request)
    await notify(lv["employee_id"], "Cuti Dibatalkan", f"Cuti {lv['jenis']} {lv['jumlah_hari']} hari dibatalkan admin. Saldo dikembalikan.", "info")
    return {"ok": True}


@api_router.delete("/leaves/{leave_id}")
async def delete_leave(leave_id: str, user: dict = Depends(get_current_user)):
    lv = await db.leaves.find_one({"id": leave_id})
    if not lv:
        raise HTTPException(status_code=404, detail="Tidak ditemukan")
    if not has_role(user, "admin") and (lv["employee_id"] != user["id"] or lv["status"] != "diajukan"):
        raise HTTPException(status_code=403, detail="Hanya pengajuan berstatus diajukan milik sendiri yang dapat dihapus")
    await db.leaves.delete_one({"id": leave_id})
    return {"ok": True}


def _leave_pdf_bytes(lv: dict, emp: dict, bal: dict, counts: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=14 * mm, bottomMargin=12 * mm, leftMargin=18 * mm, rightMargin=18 * mm)
    styles = getSampleStyleSheet()
    norm = ParagraphStyle("n", parent=styles["Normal"], fontSize=8.5, leading=11)
    normb = ParagraphStyle("nb", parent=norm, fontName="Helvetica-Bold")
    right = ParagraphStyle("r", parent=norm, alignment=2)
    center = ParagraphStyle("c", parent=norm, alignment=TA_CENTER)
    title = ParagraphStyle("t", parent=styles["Normal"], alignment=TA_CENTER, fontSize=12, fontName="Helvetica-Bold")
    sec = ParagraphStyle("sec", parent=normb, fontSize=8.5, backColor=colors.HexColor("#E2E8F0"))
    small = ParagraphStyle("sm", parent=norm, fontSize=7)

    def chk(v):
        return "[ X ]" if v else "[    ]"

    tgl = datetime.now(timezone.utc).strftime("%d %B %Y")
    jenis = lv.get("jenis", "-")
    elems = []
    elems.append(Paragraph(f"Samarinda, {tgl}", right))
    elems.append(Paragraph("Yth. Wali Kota Samarinda cq<br/>Kepala UPT Puskesmas Palaran<br/>di Tempat", norm))
    elems.append(Spacer(1, 6))
    elems.append(Paragraph("FORMULIR PERMINTAAN DAN PEMBERIAN CUTI", title))
    elems.append(Spacer(1, 8))

    def section(txt):
        return Paragraph(txt, sec)

    def kv_table(rows, w1=45 * mm):
        t = Table([[Paragraph(k, norm), Paragraph(str(v), norm)] for k, v in rows], colWidths=[w1, 174 * mm - w1])
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("LEFTPADDING", (0, 0), (-1, -1), 5)]))
        return t

    # I. DATA PEGAWAI
    elems.append(section("I. DATA PEGAWAI"))
    elems.append(kv_table([
        ("Nama", emp.get("nama", "-")), ("NIP", emp.get("nip", "-")),
        ("Jabatan", emp.get("jabatan", "-")),
        ("Masa Kerja", emp.get("masa_kerja", "-") or "-"),
        ("Unit Kerja", emp.get("unit", "-")),
    ]))
    elems.append(Spacer(1, 4))

    # II. JENIS CUTI
    elems.append(section("II. JENIS CUTI YANG DIAMBIL"))
    all_jenis = ["Tahunan", "Besar", "Sakit", "Bersalin", "Alasan Penting", "Diluar Tanggungan Negara", "Bersama"]
    jc = []
    for i, j in enumerate(all_jenis, 1):
        jc.append(Paragraph(f"{chk(j == jenis)} {i}. Cuti {j}", norm))
    jt = Table([[jc[0], jc[1]], [jc[2], jc[3]], [jc[4], jc[5]], [jc[6], ""]], colWidths=[87 * mm, 87 * mm])
    jt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("LEFTPADDING", (0, 0), (-1, -1), 5)]))
    elems.append(jt)
    elems.append(Spacer(1, 4))

    # III. ALASAN CUTI + IV. LAMANYA
    elems.append(section("III. ALASAN CUTI"))
    elems.append(kv_table([("Alasan", lv.get("alasan") or "-")]))
    elems.append(Spacer(1, 4))
    elems.append(section("IV. LAMANYA CUTI"))
    elems.append(kv_table([("Selama", f"{lv.get('jumlah_hari', 0)} Hari"),
                           ("Mulai tanggal", f"{lv.get('tanggal_mulai','-')}  s/d  {lv.get('tanggal_selesai','-')}")]))
    elems.append(Spacer(1, 4))

    # V. CATATAN CUTI TAHUNAN (SISA SALDO)
    elems.append(section("V. CATATAN CUTI TAHUNAN (SISA SALDO)"))
    y = bal.get("tahun", datetime.now(timezone.utc).year)
    saldo_rows = [
        [Paragraph("<b>Tahun</b>", small), Paragraph("<b>Sisa</b>", small), Paragraph("<b>Keterangan</b>", small)],
        [Paragraph(f"N-2 ({y-2})", small), Paragraph(str(bal.get("saldo_n2", 0)), small), Paragraph("Sisa 2 tahun sebelumnya", small)],
        [Paragraph(f"N-1 ({y-1})", small), Paragraph(str(bal.get("saldo_n1", 0)), small), Paragraph("Sisa 1 tahun sebelumnya", small)],
        [Paragraph(f"N ({y})", small), Paragraph(str(bal.get("saldo_n", 0)), small), Paragraph("Sisa tahun berjalan", small)],
        [Paragraph("Cuti Bersama", small), Paragraph(str(bal.get("saldo_bersama", 0)), small), Paragraph("Kuota cuti bersama", small)],
    ]
    left_t = Table(saldo_rows, colWidths=[26 * mm, 16 * mm, 45 * mm])
    left_t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
                                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
                                ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("LEFTPADDING", (0, 0), (-1, -1), 4)]))
    cnt_list = [("Cuti Tahunan", counts.get("Tahunan", 0)), ("Cuti Besar", 0), ("Cuti Sakit", counts.get("Sakit", 0)),
                ("Cuti Melahirkan", 0), ("Cuti Karena Alasan Penting", 0), ("Cuti Diluar Tanggungan Negara", 0), ("Cuti Bersama", 0)]
    cnt_rows = [[Paragraph(f"{i}. {n}", small), Paragraph(f"{v} kali" if v else "-", small)] for i, (n, v) in enumerate(cnt_list, 1)]
    right_t = Table(cnt_rows, colWidths=[58 * mm, 20 * mm])
    right_t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("LEFTPADDING", (0, 0), (-1, -1), 4)]))
    wrap = Table([[left_t, right_t]], colWidths=[87 * mm, 87 * mm])
    wrap.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    elems.append(wrap)
    elems.append(Spacer(1, 4))

    # VI. ALAMAT + tanda tangan pemohon
    elems.append(section("VI. ALAMAT SELAMA MENJALANKAN CUTI"))
    addr = Table([[Paragraph(f"Alamat: {lv.get('alamat') or '-'}<br/>Telp: {lv.get('telp') or '-'}", norm),
                   Paragraph(f"Hormat saya,<br/><br/><br/><br/>{emp.get('nama','-')}<br/>NIP. {emp.get('nip','-')}", center)]],
                 colWidths=[104 * mm, 70 * mm])
    addr.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 5)]))
    elems.append(addr)
    elems.append(Spacer(1, 4))

    approved = lv.get("status") == "disetujui"
    rejected = lv.get("status") == "ditolak"
    def approval_block(title_txt, checked_approve):
        opts = Paragraph(f"{chk(checked_approve)} DISETUJUI &nbsp;&nbsp; {chk(False)} PERUBAHAN &nbsp;&nbsp; {chk(False)} DITANGGUHKAN &nbsp;&nbsp; {chk(rejected)} TIDAK DISETUJUI", norm)
        sign = Paragraph("<br/><br/><br/>( ............................................ )<br/>NIP. ...............................", center)
        t = Table([[opts], [sign]], colWidths=[174 * mm])
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 5)]))
        return t
    elems.append(section("VII. PERTIMBANGAN ATASAN LANGSUNG"))
    elems.append(approval_block("atasan", approved))
    elems.append(Spacer(1, 4))
    elems.append(section("VIII. KEPUTUSAN PEJABAT YANG BERWENANG MEMBERIKAN CUTI"))
    elems.append(approval_block("pejabat", approved))
    doc.build(elems)
    return buf.getvalue()


@api_router.get("/leaves/{leave_id}/pdf")
async def leave_pdf(leave_id: str, request: Request, auth: Optional[str] = Query(None)):
    header = request.headers.get("Authorization", "")
    token = header[7:] if header.startswith("Bearer ") else auth
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")
    lv = await db.leaves.find_one({"id": leave_id})
    if not lv:
        raise HTTPException(status_code=404, detail="Tidak ditemukan")
    emp = await db.users.find_one({"id": lv["employee_id"]}) or {}
    bal = await get_balance(lv["employee_id"])
    approved = await db.leaves.find({"employee_id": lv["employee_id"], "status": "disetujui"}).to_list(2000)
    counts = {}
    for a in approved:
        counts[a["jenis"]] = counts.get(a["jenis"], 0) + 1
    return Response(content=_leave_pdf_bytes(lv, emp, bal, counts), media_type="application/pdf",
                    headers={"Content-Disposition": f"inline; filename=formulir-cuti-{leave_id[:8]}.pdf"})


@api_router.post("/leaves/{leave_id}/attachment")
async def upload_leave_attachment(leave_id: str, request: Request, file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    lv = await db.leaves.find_one({"id": leave_id})
    if not lv:
        raise HTTPException(status_code=404, detail="Pengajuan cuti tidak ditemukan")
    if not has_role(user, "admin") and lv["employee_id"] != user["id"]:
        raise HTTPException(status_code=403, detail="Akses ditolak")
    if lv["status"] != "diajukan" and not has_role(user, "admin"):
        raise HTTPException(status_code=400, detail="Lampiran hanya dapat diunggah sebelum cuti diverifikasi")
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "").lower()
    if ext not in ALLOWED_EXT:
        raise HTTPException(status_code=400, detail="Format file tidak diizinkan. Gunakan PDF, JPG, JPEG, atau PNG.")
    data = await file.read()
    if len(data) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="Ukuran file melebihi 2 MB.")
    content_type = MIME_TYPES.get(ext, "application/octet-stream")
    path = f"{APP_NAME}/leaves/{lv['employee_id']}/{uuid.uuid4()}.{ext}"
    result = put_object(path, data, content_type)
    await db.leaves.update_one({"id": leave_id}, {"$set": {
        "lampiran_path": result["path"], "lampiran_filename": file.filename, "lampiran_content_type": content_type}})
    await log_activity(user, f"Upload lampiran cuti {leave_id[:8]}", "Cuti", request)
    return {"ok": True, "lampiran_filename": file.filename}


@api_router.get("/leaves/{leave_id}/attachment")
async def download_leave_attachment(leave_id: str, request: Request, auth: Optional[str] = Query(None)):
    header = request.headers.get("Authorization", "")
    token = header[7:] if header.startswith("Bearer ") else auth
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")
    lv = await db.leaves.find_one({"id": leave_id})
    if not lv or not lv.get("lampiran_path"):
        raise HTTPException(status_code=404, detail="Lampiran tidak ditemukan")
    data, ct = get_object(lv["lampiran_path"])
    return Response(content=data, media_type=lv.get("lampiran_content_type", ct))


# ----------------------------------------------------------------------------
# SIP (Surat Izin Praktik) — input pegawai, verifikasi admin, monitoring
# ----------------------------------------------------------------------------
class SipInput(BaseModel):
    employee_id: Optional[str] = None
    nomor_sip: str
    profesi: str
    tanggal_terbit: str
    tanggal_berakhir: str
    keterangan: Optional[str] = ""


class SipVerify(BaseModel):
    status: str
    catatan: Optional[str] = ""


def sip_masa_status(tanggal_berakhir: str):
    try:
        end = datetime.fromisoformat(tanggal_berakhir[:10]).date()
    except Exception:
        return "tidak_valid", None
    days = (end - datetime.now(timezone.utc).date()).days
    if days < 0:
        return "kadaluarsa", days
    if days <= 90:
        return "akan_habis", days
    return "aktif", days


def _sip_decorate(s: dict, emp: dict):
    s = clean(s)
    ms, days = sip_masa_status(s.get("tanggal_berakhir", ""))
    s["masa_status"] = ms
    s["sisa_hari"] = days
    s["employee_nama"] = emp.get("nama", "-")
    s["employee_nip"] = emp.get("nip", "-")
    s["employee_unit"] = emp.get("unit", "-")
    return s


@api_router.get("/sip/dashboard")
async def sip_dashboard(user: dict = Depends(require_roles("admin", "kepala"))):
    sips = await db.sip.find().to_list(5000)
    ids = list({s["employee_id"] for s in sips})
    users = await db.users.find({"id": {"$in": ids}}).to_list(1000)
    umap = {u["id"]: u for u in users}
    items = [_sip_decorate(s, umap.get(s["employee_id"], {})) for s in sips]
    summary = {"total": len(items), "aktif": 0, "akan_habis": 0, "kadaluarsa": 0,
               "menunggu": sum(1 for s in items if s["status_verifikasi"] == "menunggu")}
    per_profesi = {}
    for s in items:
        if s["masa_status"] in summary:
            summary[s["masa_status"]] += 1
        p = s.get("profesi", "-")
        per_profesi[p] = per_profesi.get(p, 0) + 1
    items.sort(key=lambda x: (x["sisa_hari"] if x["sisa_hari"] is not None else 99999))
    return {"summary": summary, "per_profesi": [{"profesi": k, "jumlah": v} for k, v in per_profesi.items()], "items": items}


@api_router.get("/sip")
async def list_sip(status: Optional[str] = None, user: dict = Depends(get_current_user)):
    q = {}
    if not has_role(user, "admin", "kepala"):
        q["employee_id"] = user["id"]
    sips = await db.sip.find(q).sort("created_at", -1).to_list(2000)
    ids = list({s["employee_id"] for s in sips})
    users = await db.users.find({"id": {"$in": ids}}).to_list(1000)
    umap = {u["id"]: u for u in users}
    out = [_sip_decorate(s, umap.get(s["employee_id"], {})) for s in sips]
    if status:
        out = [s for s in out if s["masa_status"] == status]
    return out


@api_router.post("/sip")
async def create_sip(data: SipInput, request: Request, user: dict = Depends(get_current_user)):
    is_admin = has_role(user, "admin")
    emp_id = data.employee_id if (is_admin and data.employee_id) else user["id"]
    doc = {
        "id": str(uuid.uuid4()), "employee_id": emp_id, "nomor_sip": data.nomor_sip,
        "profesi": data.profesi, "tanggal_terbit": data.tanggal_terbit,
        "tanggal_berakhir": data.tanggal_berakhir, "keterangan": data.keterangan or "",
        "status_verifikasi": "disetujui" if is_admin else "menunggu",
        "catatan_verifikasi": "", "verified_by": user["nama"] if is_admin else "",
        "verified_at": now_iso() if is_admin else "", "created_at": now_iso(),
    }
    await db.sip.insert_one(dict(doc))
    await log_activity(user, f"Input SIP {data.nomor_sip}", "SIP", request)
    if not is_admin:
        for a in await db.users.find({"roles": "admin"}).to_list(100):
            await notify(a["id"], "SIP Menunggu Verifikasi", f"{user['nama']} menambahkan SIP {data.profesi}.", "info")
    return clean(doc)


@api_router.put("/sip/{sip_id}")
async def update_sip(sip_id: str, data: SipInput, user: dict = Depends(get_current_user)):
    s = await db.sip.find_one({"id": sip_id})
    if not s:
        raise HTTPException(status_code=404, detail="Tidak ditemukan")
    if not has_role(user, "admin") and s["employee_id"] != user["id"]:
        raise HTTPException(status_code=403, detail="Akses ditolak")
    upd = {"nomor_sip": data.nomor_sip, "profesi": data.profesi, "tanggal_terbit": data.tanggal_terbit,
           "tanggal_berakhir": data.tanggal_berakhir, "keterangan": data.keterangan or "", "alert_180_sent": False}
    if not has_role(user, "admin"):
        upd["status_verifikasi"] = "menunggu"
    await db.sip.update_one({"id": sip_id}, {"$set": upd})
    return {"ok": True}


@api_router.put("/sip/{sip_id}/verify")
async def verify_sip(sip_id: str, data: SipVerify, request: Request, user: dict = Depends(require_roles("admin"))):
    if data.status not in ("disetujui", "ditolak"):
        raise HTTPException(status_code=400, detail="Status tidak valid")
    if data.status == "ditolak" and not (data.catatan or "").strip():
        raise HTTPException(status_code=400, detail="Alasan penolakan wajib diisi")
    s = await db.sip.find_one({"id": sip_id})
    if not s:
        raise HTTPException(status_code=404, detail="Tidak ditemukan")
    await db.sip.update_one({"id": sip_id}, {"$set": {"status_verifikasi": data.status, "catatan_verifikasi": data.catatan or "", "verified_by": user["nama"], "verified_at": now_iso()}})
    await log_activity(user, f"Verifikasi SIP -> {data.status}", "SIP", request)
    tipe = "success" if data.status == "disetujui" else "error"
    await notify(s["employee_id"], f"SIP {data.status.capitalize()}", f"SIP {s['profesi']} Anda {data.status}.", tipe)
    return {"ok": True}


@api_router.delete("/sip/{sip_id}")
async def delete_sip(sip_id: str, user: dict = Depends(get_current_user)):
    s = await db.sip.find_one({"id": sip_id})
    if not s:
        raise HTTPException(status_code=404, detail="Tidak ditemukan")
    if not has_role(user, "admin") and s["employee_id"] != user["id"]:
        raise HTTPException(status_code=403, detail="Akses ditolak")
    await db.sip.delete_one({"id": sip_id})
    return {"ok": True}


# ----------------------------------------------------------------------------
# Include router + middleware
# ----------------------------------------------------------------------------
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup():
    try:
        init_storage()
        logger.info("Storage initialized")
    except Exception as e:
        logger.error(f"Storage init failed: {e}")
    await get_settings()
    # Migrasi multi-role: pastikan setiap pengguna punya field roles & is_blud
    async for u in db.users.find({"roles": {"$exists": False}}):
        await db.users.update_one({"id": u["id"]}, {"$set": {
            "roles": [u.get("role", "pegawai")], "is_blud": u.get("is_blud", False)}})
    async for u in db.users.find({"tipe_pegawai": {"$exists": False}}):
        await db.users.update_one({"id": u["id"]}, {"$set": {"tipe_pegawai": "BLUD" if u.get("is_blud") else "ASN"}})
    from seed import seed_data
    await seed_data(db, hash_password)
    try:
        await run_sip_alerts()
    except Exception as e:
        logger.error(f"SIP alert scan failed: {e}")


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
